from __future__ import annotations
import json, math, os, re, sqlite3, uuid
from pathlib import Path
from dotenv import load_dotenv
from flask import Flask, jsonify, render_template, request
from google import genai
from google.genai import types
from werkzeug.utils import secure_filename

load_dotenv()
ROOT = Path(__file__).resolve().parent
DATA_DIR = Path(os.getenv('DATA_DIR', ROOT / 'data'))
UPLOAD_DIR = DATA_DIR / 'uploads'
DB_PATH = DATA_DIR / 'rag.sqlite3'
ALLOWED = {'pdf', 'docx', 'txt', 'md'}
CHUNK_SIZE, OVERLAP, TOP_K = 1000, 150, 5
MODEL = os.getenv('GEMINI_MODEL', 'gemini-2.5-flash')
EMBED_MODEL = os.getenv('GEMINI_EMBEDDING_MODEL', 'gemini-embedding-001')
app = Flask(__name__)
app.config['MAX_CONTENT_LENGTH'] = 20 * 1024 * 1024
UPLOAD_DIR.mkdir(parents=True, exist_ok=True)

def connect():
    db = sqlite3.connect(DB_PATH)
    db.row_factory = sqlite3.Row
    db.execute('PRAGMA foreign_keys=ON')
    return db

def init_db():
    with connect() as db:
        db.executescript('''CREATE TABLE IF NOT EXISTS documents(id TEXT PRIMARY KEY,name TEXT NOT NULL,summary TEXT NOT NULL,created_at TEXT DEFAULT CURRENT_TIMESTAMP);
        CREATE TABLE IF NOT EXISTS chunks(id INTEGER PRIMARY KEY,document_id TEXT NOT NULL REFERENCES documents(id) ON DELETE CASCADE,chunk_index INTEGER,page INTEGER,content TEXT NOT NULL,embedding TEXT NOT NULL);
        CREATE INDEX IF NOT EXISTS chunks_document_id ON chunks(document_id);''')
init_db()

def client():
    key = os.getenv('GEMINI_API_KEY')
    if not key: raise RuntimeError('Set GEMINI_API_KEY in your .env file first.')
    return genai.Client(api_key=key)

def extract_pages(path, ext):
    if ext == 'pdf':
        import fitz
        with fitz.open(path) as pdf: return [(i, page.get_text('text')) for i, page in enumerate(pdf, 1)]
    if ext == 'docx':
        from docx import Document
        doc = Document(path)
        lines = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            lines.extend(' | '.join(c.text.strip() for c in row.cells) for row in table.rows)
        return [(None, '\n'.join(lines))]
    return [(None, path.read_text(encoding='utf-8', errors='replace'))]

def split_chunks(pages):
    result, step = [], CHUNK_SIZE - OVERLAP
    for page, text in pages:
        text = re.sub(r'\s+', ' ', text).strip()
        start = 0
        while start < len(text):
            end = min(start + CHUNK_SIZE, len(text))
            if end < len(text):
                boundary = max(text.rfind('. ', start + CHUNK_SIZE // 2, end), text.rfind(' ', start + CHUNK_SIZE // 2, end))
                if boundary > start: end = boundary + 1
            value = text[start:end].strip()
            if value: result.append((page, value))
            if end >= len(text): break
            start = max(start + step, end - OVERLAP)
    return result

def embed(api, texts, task_type):
    response = api.models.embed_content(model=EMBED_MODEL, contents=texts, config=types.EmbedContentConfig(task_type=task_type))
    return [item.values for item in response.embeddings]

def cosine(a, b):
    dot = sum(x*y for x,y in zip(a,b)); na = math.sqrt(sum(x*x for x in a)); nb = math.sqrt(sum(y*y for y in b))
    return dot/(na*nb) if na and nb else 0.0

def list_documents():
    with connect() as db:
        rows = db.execute('SELECT d.id,d.name,d.summary,d.created_at,COUNT(c.id) chunk_count FROM documents d LEFT JOIN chunks c ON c.document_id=d.id GROUP BY d.id ORDER BY d.created_at DESC').fetchall()
    return [dict(row) for row in rows]

@app.get('/')
def index(): return render_template('index.html', documents=list_documents())

@app.post('/api/upload')
def upload():
    file = request.files.get('file')
    if not file or not file.filename: return jsonify(error='Choose a document to upload.'), 400
    name = secure_filename(file.filename); ext = Path(name).suffix.lower().lstrip('.')
    if ext not in ALLOWED: return jsonify(error='Supported formats: PDF, DOCX, TXT, and Markdown.'), 400
    doc_id = uuid.uuid4().hex; path = UPLOAD_DIR / f'{doc_id}.{ext}'
    try:
        file.save(path); chunks = split_chunks(extract_pages(path, ext))
        if not chunks: return jsonify(error='No readable text found in this document.'), 400
        api = client(); vectors = []
        for offset in range(0, len(chunks), 64): vectors.extend(embed(api, [text for _,text in chunks[offset:offset+64]], 'RETRIEVAL_DOCUMENT'))
        body = '\n\n'.join(text for _,text in chunks)[:100000]
        result = api.models.generate_content(model=MODEL, contents='Summarize this document in clear, useful bullet points. Keep key names, dates, numbers, and conclusions.\n\n'+body, config=types.GenerateContentConfig(temperature=0.2,max_output_tokens=700))
        summary = (result.text or 'Summary unavailable.').strip()
        with connect() as db:
            db.execute('INSERT INTO documents(id,name,summary) VALUES(?,?,?)',(doc_id,name,summary))
            db.executemany('INSERT INTO chunks(document_id,chunk_index,page,content,embedding) VALUES(?,?,?,?,?)',[(doc_id,i,page,text,json.dumps(vector)) for i,((page,text),vector) in enumerate(zip(chunks,vectors))])
        return jsonify(document={'id':doc_id,'name':name,'summary':summary,'chunk_count':len(chunks)})
    except Exception as exc:
        app.logger.exception('Document ingestion failed')
        return jsonify(error=str(exc)),500
    finally: path.unlink(missing_ok=True)

@app.post('/api/ask')
def ask():
    data=request.get_json(silent=True) or {}; question=str(data.get('question','')).strip(); doc_id=data.get('document_id') or None
    if not question: return jsonify(error='Enter a question.'),400
    with connect() as db:
        if doc_id: rows=db.execute('SELECT c.content,c.page,c.embedding,d.name FROM chunks c JOIN documents d ON d.id=c.document_id WHERE d.id=?',(doc_id,)).fetchall()
        else: rows=db.execute('SELECT c.content,c.page,c.embedding,d.name FROM chunks c JOIN documents d ON d.id=c.document_id').fetchall()
    if not rows: return jsonify(error='Upload a document before asking a question.'),400
    try:
        api=client(); query=embed(api,[question],'RETRIEVAL_QUERY')[0]
        ranked=sorted(((cosine(query,json.loads(row['embedding'])),row) for row in rows),key=lambda pair:pair[0],reverse=True)[:TOP_K]
        selected=[(score,row) for score,row in ranked if score>0]
        if not selected: return jsonify(error="I couldn't find relevant text in the selected documents."),404
        excerpts=[]
        for _,row in selected:
            page=f", page {row['page']}" if row['page'] else ''
            excerpts.append(f"[Source: {row['name']}{page}]\n{row['content']}")
        prompt='Write a concise, direct answer in a formal tone using complete, grammatically correct sentences. Prefer one to three sentences; include only details needed to answer the question. Use only the supplied document excerpts. If they do not contain the answer, say clearly that the documents do not provide it. Cite the source document and page when available.\n\nExcerpts:\n'+'\n\n'.join(excerpts)+f'\n\nQuestion: {question}'
        answer=api.models.generate_content(model=MODEL,contents=prompt,config=types.GenerateContentConfig(temperature=0.2,max_output_tokens=800))
        sources=[{'name':row['name'],'page':row['page'],'excerpt':row['content'][:280],'score':round(score,3)} for score,row in selected]
        return jsonify(answer=(answer.text or 'No answer returned.').strip(),sources=sources)
    except Exception as exc:
        app.logger.exception('Question answering failed'); return jsonify(error=str(exc)),500

@app.get('/api/documents/<doc_id>')
def get_document(doc_id):
    with connect() as db: row=db.execute('SELECT id,name,summary FROM documents WHERE id=?',(doc_id,)).fetchone()
    if not row: return jsonify(error='Document not found.'),404
    return jsonify(dict(row))

@app.delete('/api/documents/<doc_id>')
def delete_document(doc_id):
    with connect() as db:
        # Embeddings live with their text passages in chunks; remove them first.
        db.execute('DELETE FROM chunks WHERE document_id=?',(doc_id,))
        cursor=db.execute('DELETE FROM documents WHERE id=?',(doc_id,))
    if not cursor.rowcount: return jsonify(error='Document not found.'),404
    return jsonify(ok=True)

@app.errorhandler(413)
def too_large(_error): return jsonify(error='File is too large. The limit is 20 MB.'),413

if __name__=='__main__': app.run(debug=os.getenv('FLASK_DEBUG','false').lower()=='true')
